from pathlib import Path
from datetime import datetime, date
import sys
from openpyxl import load_workbook

# Ordem compatível com as dependências do schema D1.
TABLES = [
    'CONFIGURACOES', 'DESTINATARIOS', 'USUARIOS', 'AUDIENCIAS',
    'AUDIENCIA_DESTINATARIOS', 'OFICIOS', 'NOTIFICACOES', 'EVENTOS', 'AUDITORIA'
]

# Colunas aceitas pelo schema D1 atual. Campos legados que existam na planilha,
# mas não no D1 (ex.: biometria antiga), são ignorados de forma explícita.
D1_COLUMNS = {
    'AUDIENCIAS': {'id','codigo','processo','assunto','destinatario_id','data_hora','timezone','local','modalidade','link','observacoes','oficio_drive_id','status','criado_em','criado_por','atualizado_em','atualizado_por'},
    'CONFIGURACOES': {'chave','valor','atualizado_em','atualizado_por'},
    'OFICIOS': {'id','audiencia_id','destinatario_id','numero','ano','status','drive_id','data_emissao','destino_judicial','trecho_opcional','posto_graduacao','dados_json','pdf_nome','processo_pdf_drive_id','processo_pdf_nome','criado_em','criado_por','atualizado_em','enviado_em','confirmado_em','r2_pdf_key','r2_processo_key'},
    'AUDIENCIA_DESTINATARIOS': {'id','audiencia_id','destinatario_id','status','criado_em','criado_por'},
    'DESTINATARIOS': {'id','nome','telefone','identificacao','unidade','status','criado_em','atualizado_em','rg','cpf','posto_graduacao'},
    'NOTIFICACOES': {'id','chave_idempotencia','audiencia_id','destinatario_id','oficio_id','tipo','telefone','template','mensagem','token_hash','data_programada','status','tentativas','provider','wamid','ultimo_erro_codigo','ultimo_erro_resumo','criado_em','criado_por','processando_em','simulado_em','enviado_em','confirmado_em','atualizado_em'},
    'EVENTOS': {'id','audiencia_id','notificacao_id','tipo','timestamp','wamid','origem','dados'},
    'USUARIOS': {'id','nome','login','perfil','status','senha_hash','salt','ultimo_login','criado_em','atualizado_em','tipo_conta','troca_senha_pendente','tentativas_falhas','bloqueado_ate','ultimo_erro_login','ultimo_ip_contexto'},
    'AUDITORIA': {'id','timestamp','usuario_id','acao','entidade','registro_id','ip_contexto','dados_anteriores','dados_novos','resultado'},
}

LEGACY_USER_COLUMNS = {
    'biometria_credential_id', 'biometria_token_hash', 'biometria_ativada_em',
    'biometria_ultimo_uso', 'biometria_dispositivo'
}

def q(v):
    if v is None or v == '': return 'NULL'
    if isinstance(v, (datetime, date)): return "'" + v.isoformat() + "'"
    if isinstance(v, bool): return '1' if v else '0'
    if isinstance(v, (int, float)) and not isinstance(v, bool): return str(v)
    s = str(v).strip()
    return "'" + s.replace("'", "''") + "'"

def main():
    src = Path(sys.argv[1] if len(sys.argv) > 1 else 'SISTEMA_AUDIENCIAS_DB.xlsx')
    out = Path(sys.argv[2] if len(sys.argv) > 2 else '.local/import_current.sql')
    wb = load_workbook(src, data_only=True, read_only=True)
    out.parent.mkdir(parents=True, exist_ok=True)

    # Não usar BEGIN/COMMIT: wrangler d1 execute --remote rejeita transações SQL explícitas.
    lines = ['PRAGMA foreign_keys=OFF;']
    total = 0

    for sheet_name in TABLES:
        ws = wb[sheet_name]
        rows = ws.iter_rows(values_only=True)
        headers = [str(x).strip() if x is not None else '' for x in next(rows)]
        allowed = D1_COLUMNS[sheet_name]
        valid = [(i, h) for i, h in enumerate(headers) if h and h in allowed]
        ignored = [h for h in headers if h and h not in allowed]

        unexpected = [h for h in ignored if not (sheet_name == 'USUARIOS' and h in LEGACY_USER_COLUMNS)]
        if unexpected:
            raise ValueError(f'{sheet_name}: coluna(s) não reconhecida(s) pelo schema D1: {", ".join(unexpected)}')
        if ignored:
            print(f'{sheet_name}: ignorando coluna(s) legada(s): {", ".join(ignored)}')

        table = sheet_name.lower()
        count = 0
        for row in rows:
            vals = [row[i] if i < len(row) else None for i, _ in valid]
            if not any(v not in (None, '') for v in vals):
                continue
            cols = ','.join('"' + h.replace('"', '""') + '"' for _, h in valid)
            values = ','.join(q(v) for v in vals)
            lines.append(f'INSERT OR REPLACE INTO "{table}" ({cols}) VALUES ({values});')
            count += 1
            total += 1
        print(f'{sheet_name}: {count} registro(s)')

    lines.append('PRAGMA foreign_keys=ON;')
    out.write_text('\n'.join(lines) + '\n', encoding='utf-8')
    print(f'Total: {total} registro(s)')
    print(f'Arquivo: {out.resolve()}')

if __name__ == '__main__':
    main()
